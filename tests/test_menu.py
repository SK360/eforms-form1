"""The ./atf menu (src/ui.py) with scripted answers: adding, filing, resolving, and screenshot-safe output."""
import csv
import os
import sys

import pytest

from conftest import TICKED, atf, make_individual, pick, queue, ui

QUIT = lambda kw: "quit"  # noqa: E731


def test_add_then_file_two(configured, portal, answers, screen):
    queue(("T101", "22", "filed", "900000101"), ("T102", "9", "filed", "900000102"))
    portal([("T103", "9"), ("T104", "22")])
    checks = {}

    def serial(kw):
        checks["dup"], checks["blank"] = kw["validate"]("t102"), kw["validate"]("  ")
        return kw["default"]
    answers(pick("Add suppressors"), "1", serial, lambda kw: kw["default"], "6.5", True, "1", lambda kw: kw["default"], "22",
            lambda kw: (checks.setdefault("len_default", kw["default"]), kw["default"])[1], False,
            pick("File 2 queued"), TICKED, False, True, QUIT)
    ui.menu()
    out, rows = screen.getvalue(), atf.load_rows()
    assert [(r["serial_number"], r["caliber"], r["status"]) for r in rows[2:]] == [("T103", "9", "filed"), ("T104", "22", "filed")]
    assert checks == {"dup": "That serial is already in the list", "blank": "Enter a serial number", "len_default": "6.5"}
    assert [r["overall_length"] for r in rows[2:]] == ["6.5", "6.5"]                       # per suppressor, recorded
    for shown in ("✔ Applicant type", "✔ Responsible person", "Serial number", "Your turn in the browser",
                  "Filed: application #…0001", "Filed: application #…0002", "Result"):
        assert shown in out, shown
    assert "DOE FAMILY TRUST" not in out.upper() and "900000102" not in out and "(hidden)" in out  # screenshot-safe
    assert (atf.OUT / "5320.23_T103.pdf").exists()


def test_file_as_individual_from_the_menu(configured, portal, answers, screen):
    queue(("S1", "9"))
    make_individual()
    portal([("S1", "9")])
    answers(pick("File 1 queued"), TICKED, False, True, QUIT)
    ui.menu()
    out = screen.getvalue()
    assert "Filing as: an individual (you)" in out and "Questions 11.A-11.M.1: all No; 11.M.2: N/A" in out
    assert "answers to questions 11.A-11.M above are true today" in out and "5320.23s generated" not in out
    assert atf.load_rows()[0]["status"] == "filed" and "Answer 11.A" in out


def test_practice_run_never_submits(configured, portal, answers, screen):
    queue(("S1", "9"))
    state = portal([("S1", "9")])
    answers(pick("Advanced"), pick("Practice run"), lambda kw: kw["choices"][0].value, True, QUIT)
    ui.menu()
    out = screen.getvalue()
    assert atf.load_rows()[0]["status"] == "" and state["logged_out"]           # still queued: nothing signed
    assert "Your turn in the browser" not in out and "Practice run (nothing is submitted)" in out
    assert "Serial number" in out                                                # the read-back check still runs


def test_declining_confirmation_opens_no_browser(configured, answers, screen, monkeypatch):
    queue(("S1", "9"))
    monkeypatch.setattr(atf, "_browser", lambda: (_ for _ in ()).throw(AssertionError("browser opened")))
    answers(pick("File 1 queued"), TICKED, False, False, QUIT)
    ui.menu()
    assert atf.load_rows()[0]["status"] == ""
    assert answers.kw["Confirm and open the browser?"]["auto_enter"] is False   # 'y' alone doesn't start it: Enter does


@pytest.mark.skipif(os.name == "nt", reason="uses a POSIX pseudo-terminal")
def test_keys_typed_ahead_cannot_answer_the_login_prompt(monkeypatch):
    """An Enter pressed after 'y' on "Open the browser?" must not answer "Log in ... then press Enter"."""
    import pty
    import select
    master, slave = pty.openpty()
    tty = os.fdopen(slave, "r")
    os.write(master, b"\n")                                        # the stray Enter, typed before the prompt
    assert select.select([tty], [], [], 1)[0]                       # it's waiting in the terminal
    monkeypatch.setattr(sys, "stdin", tty)
    ui._drop_typeahead()
    assert not select.select([tty], [], [], 0.2)[0]                 # gone
    tty.close()
    os.close(master)


def test_incomplete_settings_block_filing(configured, answers, screen):
    queue(("S1", "9"))
    (atf.FORMS / "5320.23.pdf").unlink()
    answers(pick("File 1 queued"), QUIT)
    ui.menu()
    assert "Fix these in Settings first" in screen.getvalue() and atf.load_rows()[0]["status"] == ""


def test_resolve_pending(configured, answers, screen):
    queue(("S1", "9", "pending"), ("S2", "9", "pending"))
    checks = {}

    def number(kw):
        checks["bad"] = kw["validate"]("20-27")
        return "900000201"
    answers(pick("⚠ Resolve"), lambda kw: kw["choices"][0].value, "filed", number, pick("⚠ Resolve"), "queued", QUIT)
    ui.menu()
    rows = atf.load_rows()
    assert (rows[0]["status"], rows[0]["application_number"], rows[1]["status"]) == ("filed", "900000201", "")
    assert checks["bad"] == "Digits only"
    with open(atf.OUT / "submissions.csv", newline="", encoding="utf-8") as f:
        assert list(csv.reader(f))[-1][1:] == ["S1", "9", "900000201"]


def test_home_shows_last_five_filed_and_short_numbers(configured, answers, screen):
    queue(*[(f"T{n:03d}", "9", "filed", f"9000000{n:02d}") for n in range(1, 11)], ("T011", "9"))
    answers(QUIT)
    ui.menu()
    out = screen.getvalue()
    assert "… 5 earlier" in out and "T005" not in out and "T006" in out and "T011" in out
    assert "…0010" in out and "900000010" not in out


@pytest.mark.parametrize("last, suggested", [
    ("A046", "A047"), ("A099", "A100"), ("ABC-0099", "ABC-0100"), ("0012", "0013"), ("SN99", "SN100"),
    ("MS-2026-01", "MS-2026-02"), ("2026-001A", "2026-002A"), ("abc12x", "abc13x"), ("AB 12", "AB 13"),
    ("XYZ", ""), ("", "")])
def test_next_serial(last, suggested):
    assert ui.next_serial(last) == suggested


def test_short_application_number():
    assert (ui.short("900000102"), ui.short(""), ui.short("UNKNOWN"), ui.short("123")) == ("…0102", "", "UNKNOWN", "123")


def test_serial_key_is_natural():
    assert sorted(["S10", "S9", "s100", "S10A", "A1"], key=atf.serial_key) == ["A1", "S9", "S10", "S10A", "s100"]


def test_filed_dates_backfilled_from_the_log(configured):
    queue(("T1", "9", "filed", "900000001"), ("T2", "9"))
    (atf.OUT / "submissions.csv").write_text("timestamp,serial_number,caliber,application_number\n2026-10-08T21:57:48,T1,9,900000001\n",
                                             encoding="utf-8")
    rows = atf.backfill_filed_dates()
    assert rows[0]["filed_date"] == "10/08/2026" and rows[1]["filed_date"] == ""
    assert atf.load_rows()[0]["filed_date"] == "10/08/2026"


def test_approved_rows_are_never_refiled(configured):
    queue(("T1", "9", "approved"), ("T2", "9"))
    assert atf.pick_row(atf.load_rows(), None) == 1
    with pytest.raises(SystemExit, match="approved"):
        atf.pick_row(atf.load_rows(), 1)


def test_inventory_screen(configured, answers, screen):
    atf.save_rows([
        {"serial_number": "T10", "caliber": "22", "status": "approved", "approved_date": "08/15/2026", "eform_id": "1000010", "conditions": ""},
        {"serial_number": "T9", "caliber": "22", "status": "approved", "approved_date": "02/14/2026", "eform_id": "1000009",
         "conditions": "MARKING"},
        {"serial_number": "T11", "caliber": "9", "status": "filed", "filed_date": "10/01/2026", "application_number": "900000011"},
        {"serial_number": "T12", "caliber": "9", "status": ""}])
    answers(pick("Inventory"), "approved", "back", QUIT)
    ui.menu()
    out = screen.getvalue()
    full = out.split("Inventory · All (4)")[1].split("Inventory · Approved")[0]
    assert full.index("T9") < full.index("T10") < full.index("T11") < full.index("T12")        # serial order
    assert "waiting" in full and "…0011" in full and "900000011" not in full and "⚠ conditions" in full
    assert "4 suppressors · 2 approved · 1 in progress · 1 queued" in out
    approved = out.split("Inventory · Approved (2)")[1].split("eForms Form 1")[0]                # just that view, not the home screen after
    assert "T9" in approved and "T11" not in approved and "T12" not in approved


def _inventory_rows():
    atf.save_rows([
        {"serial_number": "T2", "caliber": "22", "overall_length": "12", "status": "approved", "approved_date": "08/15/2026",
         "eform_id": "1000002", "control_number": "20260000002"},
        {"serial_number": "T10", "caliber": "9", "overall_length": "12", "status": "filed", "filed_date": "10/08/2026",
         "application_number": "900000010"},
        {"serial_number": "T11", "caliber": "9", "overall_length": "12", "status": ""},
        {"serial_number": "T12", "caliber": "9", "overall_length": "6.5", "status": ""}])


def test_edit_and_remove_queued(configured, answers, screen):
    _inventory_rows()
    offered = {}

    def which(kw):
        offered["rows"] = [c.title for c in kw["choices"]]
        return kw["choices"][0].value                                  # T11
    dup = {}

    def new_serial(kw):
        dup["msg"] = kw["validate"]("t12")
        return "T13"
    answers(pick("Edit or remove"), which, "serial_number", new_serial,
            pick("Edit or remove"), lambda kw: kw["choices"][1].value, "overall_length", "7.25",   # T13 is 2nd now
            pick("Edit or remove"), lambda kw: kw["choices"][0].value, "remove", True,              # remove T12
            QUIT)
    ui.menu()
    rows = {r["serial_number"]: r for r in atf.load_rows()}
    assert offered["rows"][0].startswith("T11") and len(offered["rows"]) == 3 and not any("T2 " in t or "T10" in t for t in offered["rows"])
    assert dup["msg"] == "That serial is already in the list"
    assert set(rows) == {"T2", "T10", "T13"} and rows["T13"]["overall_length"] == "7.25"
    assert rows["T2"]["status"] == "approved" and rows["T10"]["status"] == "filed"              # locked rows untouched


def test_open_approved_pdf(configured, answers, screen, monkeypatch):
    _inventory_rows()
    folder = atf.OUT / "approved"
    folder.mkdir()
    (folder / "T2.pdf").write_bytes(b"%PDF")
    (folder / "index.csv").write_text("eform_id,file\n1000002,T2.pdf\n", encoding="utf-8")
    opened = []
    monkeypatch.setattr(ui, "open_file", lambda p: opened.append(p))
    answers(pick("Inventory"), pick("Open an approved PDF"), lambda kw: kw["choices"][0].value, "back", QUIT)
    ui.menu()
    assert opened == [folder / "T2.pdf"]


def test_export_inventory(configured):
    from openpyxl import load_workbook
    _inventory_rows()
    path = atf.export_inventory(atf.OUT / "inv.xlsx")
    ws = load_workbook(path).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:4] == ("Serial", "Caliber", "Length (in)", "Status")
    assert [r[0] for r in rows[1:]] == ["T2", "T10", "T11", "T12"]                             # serial order
    assert rows[4][2] == 6.5 and rows[3][3] == "queued" and ws.freeze_panes == "A2"
    assert rows[1][9] == "20260000002"                                                         # full numbers in the file


def test_add_a_run_of_serials(configured, answers, screen):
    queue(("T054", "9", "filed", "900000054"))
    seen = {}

    def first(kw):
        seen["default"], seen["letters"] = kw["default"], kw["validate"]("ABC")
        return kw["default"]
    answers(pick("Add suppressors"), "10", first, "9", "12", True, False, QUIT)
    ui.menu()
    rows = atf.load_rows()
    assert seen == {"default": "T055", "letters": "To add several, the serial needs a number to count up from"}
    assert [r["serial_number"] for r in rows[1:]] == [f"T0{n}" for n in range(55, 65)]
    assert {(r["caliber"], r["overall_length"], r["status"]) for r in rows[1:]} == {("9", "12", "")}
    assert "Added T055 – T064 (10 suppressors)" in screen.getvalue()


def test_a_run_that_would_clash_adds_nothing(configured, answers, screen):
    queue(("T054", "9", "filed"), ("T057", "9"))
    answers(pick("Add suppressors"), "5", "T055", "9", "12",          # T055..T059 would include T057
            "3", "T058", "9", "12", True, False, QUIT)                  # then a custom count that fits: T058..T060
    ui.menu()
    serials = [r["serial_number"] for r in atf.load_rows()]
    assert "Already in your list: T057. Nothing added." in screen.getvalue()
    assert serials == ["T054", "T057", "T058", "T059", "T060"]


def test_serial_run():
    assert ui.serial_run("A-0098", 3) == ["A-0098", "A-0099", "A-0100"]
    assert ui.serial_run("2026-009A", 2) == ["2026-009A", "2026-010A"]
    assert ui.serial_run("XYZ", 2) == [] and ui.serial_run("XYZ", 1) == ["XYZ"]



def _sbrs():
    atf.save_sbrs([{"serial_number": "R2", "caliber": "5.56", "model": "X15", "barrel_length": "10.5", "overall_length": "26",
                    "approved_date": "09/13/2026", "eform_id": "1000020", "control_number": "", "conditions": ""},
                   {"serial_number": "R10", "caliber": "300", "model": "Y9", "barrel_length": "8", "overall_length": "24",
                    "approved_date": "09/14/2026", "eform_id": "1000021", "control_number": "", "conditions": "MARK IT"}])


def test_sbr_inventory_view_and_summary(configured, answers, screen):
    _inventory_rows()
    _sbrs()
    answers(pick("Inventory"), "sbrs", "back", QUIT)
    ui.menu()
    out = screen.getvalue()
    view = out.split("Inventory · SBRs (2)")[1].split("eForms Form 1")[0]
    assert view.index("R2") < view.index("R10") and "10.5" in view and "X15" in view and "eForm …0021" in view
    assert "T11" not in view                                                                    # suppressors not in this view
    assert view.count("✔ approved") == 2 and "2 SBRs · 2 approved" in view                     # same look as suppressors
    assert "suppressors ·" not in view                                                          # its own summary line
    assert "⚠ approved with conditions: R10" in view and "data/sbrs.csv" in view
    assert "· 2 SBRs" in out


def test_open_sbr_pdf_and_export_sheet(configured, answers, screen, monkeypatch):
    from openpyxl import load_workbook
    _inventory_rows()
    _sbrs()
    folder = atf.OUT / "approved"
    folder.mkdir()
    (folder / "R2.pdf").write_bytes(b"%PDF")
    (folder / "index.csv").write_text("eform_id,file\n1000020,R2.pdf\n", encoding="utf-8")
    opened = []
    monkeypatch.setattr(ui, "open_file", lambda p: opened.append(p))
    answers(pick("Inventory"), pick("Open an approved PDF"), lambda kw: next(c.value for c in kw["choices"] if c.title.startswith("R2")),
            "back", QUIT)
    ui.menu()
    assert opened == [folder / "R2.pdf"]
    wb = load_workbook(atf.export_inventory(atf.OUT / "inv.xlsx"))
    rows = list(wb["SBRs"].iter_rows(values_only=True))
    assert wb.sheetnames == ["Inventory", "SBRs"] and [r[0] for r in rows[1:]] == ["R2", "R10"] and rows[1][3] == 10.5
